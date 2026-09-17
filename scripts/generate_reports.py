"""
generate_reports.py
--------------------
Generates the Excel workbook (Pediatric_Camp_Cleaned_Data.xlsx) with the
required sheets: Cleaned Data, Summary, Data Quality, Duplicate Merge Log,
Validation Report.
"""
from __future__ import annotations

import os
from typing import Any, Dict

import pandas as pd


def _summary_df(summary: Dict[str, Any]) -> pd.DataFrame:
    rows = [{"Metric": k, "Value": v} for k, v in summary.get("kpis", {}).items()]
    return pd.DataFrame(rows, columns=["Metric", "Value"])


def _validation_df(df: pd.DataFrame) -> pd.DataFrame:
    flag_cols = [c for c in df.columns if c.endswith("_Quality_Flag")]
    if not flag_cols:
        return pd.DataFrame(columns=["Field", "Valid", "Missing", "Potentially_Invalid", "Needs_Review"])
    rows = []
    for col in flag_cols:
        vc = df[col].value_counts()
        rows.append({
            "Field": col.replace("_Quality_Flag", ""),
            "Valid": int(vc.get("Valid", 0)),
            "Missing": int(vc.get("Missing", 0)),
            "Potentially_Invalid": int(vc.get("Potentially Invalid", 0)),
            "Needs_Review": int(vc.get("Needs Review", 0)),
        })
    return pd.DataFrame(rows)


def write_excel_report(
    cleaned_df: pd.DataFrame,
    summary: Dict[str, Any],
    data_quality_df: pd.DataFrame,
    merge_log_df: pd.DataFrame,
    excel_dir: str,
    filename: str = "Pediatric_Camp_Cleaned_Data.xlsx",
) -> str:
    os.makedirs(excel_dir, exist_ok=True)
    out_path = os.path.join(excel_dir, filename)

    export_cols = [c for c in cleaned_df.columns if not c.startswith("_")]

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        cleaned_df[export_cols].to_excel(writer, sheet_name="Cleaned Data", index=False)
        _summary_df(summary).to_excel(writer, sheet_name="Summary", index=False)
        data_quality_df.to_excel(writer, sheet_name="Data Quality", index=False)
        merge_log_df.to_excel(writer, sheet_name="Duplicate Merge Log", index=False)
        _validation_df(cleaned_df).to_excel(writer, sheet_name="Validation Report", index=False)

    # Light formatting pass: bold headers, autosize columns (best-effort)
    try:
        from openpyxl.styles import Font
        from openpyxl import load_workbook

        wb = load_workbook(out_path)
        for ws in wb.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for col_cells in ws.columns:
                max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
                col_letter = col_cells[0].column_letter
                ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 45)
        wb.save(out_path)
    except Exception:  # noqa: BLE001
        pass  # formatting is best-effort only

    return out_path
